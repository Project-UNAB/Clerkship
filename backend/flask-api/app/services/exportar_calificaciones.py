"""Exporta el libro de calificaciones de un curso a CSV o XLSX.

Una fila por estudiante y una columna por tarea o cuestionario, más el
promedio. Los nombres de estudiantes, tareas y cursos los escribió alguien:
si una celda empieza con =, +, - o @, una hoja de cálculo la ejecutaría como
fórmula, así que esas celdas se guardan como texto.
"""
import csv
import io
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATOS = ("csv", "xlsx")
MIME = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_INICIO_DE_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def texto_seguro(valor) -> str:
    """El texto tal cual, salvo que una hoja de cálculo lo fuera a tomar por
    fórmula: en ese caso se le antepone un apóstrofo (queda como texto)."""
    texto = "" if valor is None else str(valor)
    return "'" + texto if texto.startswith(_INICIO_DE_FORMULA) else texto


def nombre_de_archivo(nombre_curso: str, formato: str, fecha) -> str:
    """calificaciones-gastroenterologia-i-2026-10-10.xlsx — solo letras, números y guiones."""
    plano = unicodedata.normalize("NFKD", nombre_curso or "").encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-z0-9]+", "-", plano.lower()).strip("-")[:60] or "curso"
    return f"calificaciones-{slug}-{fecha:%Y-%m-%d}.{formato}"


def construir_tabla(items_meta, estudiantes) -> tuple:
    """(encabezados, filas). Las notas van como números; lo entregado sin
    calificar como "Sin calificar" y lo no entregado, vacío."""
    encabezados = ["Código", "Estudiante", "Correo"]
    for item in items_meta:
        tipo = "Tarea" if item["type"] == "ASSIGNMENT" else "Cuestionario"
        maximo = f" (máx. {item['max_score']:g})" if item.get("max_score") is not None else ""
        encabezados.append(texto_seguro(f"{tipo}: {item['title']}{maximo}"))
    encabezados.append("Promedio (%)")

    filas = []
    for estudiante in estudiantes:
        fila = [
            texto_seguro(estudiante.get("student_code")),
            texto_seguro(estudiante.get("nombre")),
            texto_seguro(estudiante.get("email")),
        ]
        for item in items_meta:
            celda = (estudiante.get("calificaciones") or {}).get(item["id"])
            if celda is None:
                fila.append("")
            elif celda.get("score") is None:
                fila.append("Sin calificar")
            else:
                fila.append(float(celda["score"]))
        fila.append(estudiante.get("promedio") if estudiante.get("promedio") is not None else "")
        filas.append(fila)
    return encabezados, filas


def a_csv(encabezados, filas) -> bytes:
    """CSV en UTF-8 con BOM (para que Excel lea bien tildes y eñes)."""
    salida = io.StringIO(newline="")
    escritor = csv.writer(salida, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    escritor.writerow(encabezados)
    escritor.writerows(filas)
    return ("﻿" + salida.getvalue()).encode("utf-8")


def a_xlsx(encabezados, filas, titulo_hoja="Calificaciones") -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = re.sub(r"[\[\]\*\?/\\:]", " ", titulo_hoja)[:31] or "Calificaciones"

    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0369A1")
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        hoja.append(fila)

    # Ninguna celda de texto queda como fórmula, pase lo que pase con su contenido.
    for fila in hoja.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str):
                celda.data_type = "s"

    hoja.freeze_panes = "C2"   # encabezado y nombre siempre a la vista
    for indice, encabezado in enumerate(encabezados, start=1):
        ancho = 28 if indice == 2 else (30 if indice == 3 else max(12, min(32, len(str(encabezado)) + 2)))
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.row_dimensions[1].height = 34

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def exportar(formato, items_meta, estudiantes, titulo_hoja="Calificaciones") -> bytes:
    encabezados, filas = construir_tabla(items_meta, estudiantes)
    if formato == "xlsx":
        return a_xlsx(encabezados, filas, titulo_hoja)
    return a_csv(encabezados, filas)
