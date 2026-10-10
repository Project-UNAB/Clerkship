"""Fase 7 — funcionalidades adicionales: notificaciones (in-app y por correo),
calificación con rúbrica, exportar el libro de calificaciones a CSV y XLSX, e
indicadores de la card del curso para el estudiante.

No toca la base real, R2 ni Mailgun — todo con monkeypatch."""
import csv
import io
import json
import uuid
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest
from flask_jwt_extended import create_access_token
from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql

import app.routes.curso_calificaciones as cal_routes
import app.routes.curso_contenido as cc_routes
import app.routes.cursos as cursos_routes
import app.routes.notificaciones as noti_routes
from app import mailer
from app.email_templates import notification_email_html
from app.models import AssignmentSubmission, AuditLog, Course, CourseContentItem, Notification
from app.services import exportar_calificaciones, notificaciones, tareas
# Las funciones reales: conftest las reemplaza en el módulo para el resto de la suite.
from app.services.notificaciones import aviso_publicado, tarea_calificada, tarea_publicada
from tests.fakes import ConsultaFalsa

AHORA = datetime.now(timezone.utc)


def _headers(app, role, identity):
    with app.app_context():
        token = create_access_token(identity=str(identity), additional_claims={"role": role})
    return {"Authorization": f"Bearer {token}"}


# ═══════════════════════════ 1. Notificaciones ═══════════════════════════

class _Avisos:
    """Captura lo que el servicio de notificaciones escribiría y enviaría."""

    def __init__(self, app, monkeypatch, estudiantes, con_correo=()):
        self.creadas = []          # (user_ids, tipo, titulo, cuerpo, datos)
        self.correos = []          # (destinatarios, asunto, texto)
        self.commits = 0
        self.ya_notificados = set()
        yo = self

        def crear(user_ids, tipo, titulo, cuerpo=None, dedupe=None, **datos):
            nuevos = []
            for uid in user_ids:
                clave = dedupe(uid) if dedupe else None
                if clave and clave in yo.ya_notificados:
                    continue
                if clave:
                    yo.ya_notificados.add(clave)
                nuevos.append((uuid.uuid4(), uid))
            yo.creadas.append(([uid for _n, uid in nuevos], tipo, titulo, cuerpo, datos))
            return nuevos

        def enviar_correos(creadas, asunto, texto):
            destinatarios = [uid for _n, uid in creadas if uid in con_correo]
            if destinatarios:
                yo.correos.append((destinatarios, asunto, texto))
            return len(destinatarios)

        with app.app_context():
            monkeypatch.setattr(notificaciones, "_estudiantes_del_curso", lambda course_id: list(estudiantes))
            monkeypatch.setattr(notificaciones, "_crear", crear)
            monkeypatch.setattr(notificaciones, "_enviar_correos", enviar_correos)
            monkeypatch.setattr(notificaciones.db.session, "commit", lambda: setattr(yo, "commits", yo.commits + 1))
            monkeypatch.setattr(notificaciones.db.session, "rollback", lambda: None)


def _curso_y_tarea(**kw):
    curso = SimpleNamespace(id=uuid.uuid4(), name="Gastro I", teacher_id=uuid.uuid4())
    datos = dict(id=uuid.uuid4(), title="Informe de caso", due_at=None, max_score=10, type="ASSIGNMENT")
    datos.update(kw)
    return curso, SimpleNamespace(**datos)


def test_nueva_tarea_avisa_a_todos_los_estudiantes_del_curso(app, monkeypatch):
    a, b = uuid.uuid4(), uuid.uuid4()
    avisos = _Avisos(app, monkeypatch, [a, b], con_correo=[b])
    cierre = datetime(2026, 10, 28, 4, 59, tzinfo=timezone.utc)
    curso, tarea = _curso_y_tarea(due_at=cierre)

    with app.app_context():
        assert tarea_publicada(curso, tarea) == 2

    user_ids, tipo, titulo, cuerpo, datos = avisos.creadas[0]
    assert user_ids == [a, b] and tipo == "ASSIGNMENT_PUBLISHED"
    assert titulo == "Nueva tarea: Informe de caso" and "Gastro I" in cuerpo
    assert datos == {"course_id": curso.id, "entity_type": "assignment", "entity_id": tarea.id, "event_at": cierre}
    # Por correo, solo a quien lo activó; y la fecha va escrita en hora de Colombia.
    destinatarios, asunto, texto = avisos.correos[0]
    assert destinatarios == [b] and "Informe de caso" in asunto
    assert "27 oct 2026, 11:59 p. m. (hora de Colombia)" in texto


def test_nuevo_aviso_y_tarea_calificada(app, monkeypatch):
    a, b = uuid.uuid4(), uuid.uuid4()
    avisos = _Avisos(app, monkeypatch, [a, b])
    curso, tarea = _curso_y_tarea()
    aviso = SimpleNamespace(id=uuid.uuid4(), title="Cambio de salón")
    entrega = SimpleNamespace(student_id=b, score=8.5)

    with app.app_context():
        assert aviso_publicado(curso, aviso) == 2
        assert tarea_calificada(curso, tarea, entrega) == 1

    (_, tipo_aviso, titulo_aviso, _, datos_aviso), (para, tipo_nota, titulo_nota, cuerpo_nota, _) = avisos.creadas
    assert tipo_aviso == "ANNOUNCEMENT" and titulo_aviso == "Nuevo aviso: Cambio de salón"
    assert datos_aviso["entity_type"] == "announcement" and datos_aviso["entity_id"] == aviso.id
    # La nota le llega solo al estudiante calificado.
    assert para == [b] and tipo_nota == "ASSIGNMENT_GRADED"
    assert titulo_nota == "Tarea calificada: Informe de caso" and "8.5 / 10" in cuerpo_nota


def test_si_avisar_falla_la_accion_principal_no_se_entera(app, monkeypatch):
    curso, tarea = _curso_y_tarea()
    rollbacks = []

    def roto(*a, **kw):
        raise RuntimeError("relation notifications does not exist")

    with app.app_context():
        monkeypatch.setattr(notificaciones, "_estudiantes_del_curso", lambda course_id: [uuid.uuid4()])
        monkeypatch.setattr(notificaciones, "_crear", roto)
        monkeypatch.setattr(notificaciones.db.session, "rollback", lambda: rollbacks.append(True))
        assert tarea_publicada(curso, tarea) == 0          # no lanza
    assert rollbacks == [True]


def test_si_el_correo_falla_la_notificacion_in_app_queda(app, monkeypatch):
    a = uuid.uuid4()
    avisos = _Avisos(app, monkeypatch, [a])
    curso, tarea = _curso_y_tarea()

    def mailgun_caido(*args, **kw):
        raise RuntimeError("503 Service Unavailable")

    with app.app_context():
        monkeypatch.setattr(notificaciones, "_enviar_correos", mailgun_caido)
        assert tarea_publicada(curso, tarea) == 1
    assert avisos.commits == 1          # la notificación in-app ya estaba confirmada


def test_crear_inserta_una_fila_por_usuario_e_ignora_repetidos(app, monkeypatch):
    ejecutadas = []
    a, b = uuid.uuid4(), uuid.uuid4()

    def execute(stmt):
        ejecutadas.append(stmt)
        return SimpleNamespace(all=lambda: [(uuid.uuid4(), a)])

    with app.app_context():
        monkeypatch.setattr(notificaciones.db.session, "execute", execute)
        creadas = notificaciones._crear(
            [a, b, a, None], "DUE_SOON", "Cierra pronto: X", "cuerpo", dedupe=lambda uid: f"due:item:{uid}",
        )
        assert notificaciones._crear([], "DUE_SOON", "x") == []

    assert [uid for _n, uid in creadas] == [a]          # lo que la base dice que insertó
    compilada = ejecutadas[0].compile(dialect=postgresql.dialect())
    assert "INSERT INTO notifications" in str(compilada)
    assert "ON CONFLICT (dedupe_key) DO NOTHING" in str(compilada)
    claves = [v for k, v in compilada.params.items() if k.startswith("dedupe_key")]
    assert sorted(claves) == sorted([f"due:item:{a}", f"due:item:{b}"])          # sin repetidos ni vacíos


def test_recordatorio_de_24_horas_solo_a_quien_no_ha_entregado_y_una_sola_vez(app, monkeypatch):
    entrego, falta_1, falta_2 = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    avisos = _Avisos(app, monkeypatch, [entrego, falta_1, falta_2], con_correo=[falta_1])
    curso = SimpleNamespace(id=uuid.uuid4(), name="Gastro I")
    cierre = AHORA + timedelta(hours=20)
    tarea = SimpleNamespace(id=uuid.uuid4(), title="Informe", due_at=cierre)

    def consulta(*columnas):
        if columnas[0] is notificaciones.CourseContentItem:
            return ConsultaFalsa(todos=[(tarea, curso)])
        return ConsultaFalsa(todos=[(entrego,)])          # quién ya entregó

    with app.app_context():
        monkeypatch.setattr(notificaciones.db.session, "query", consulta)
        primero = notificaciones.recordar_vencimientos(AHORA)
        segundo = notificaciones.recordar_vencimientos(AHORA + timedelta(hours=1))

    assert primero == {"tareas": 1, "notificaciones": 2}
    user_ids, tipo, titulo, _cuerpo, datos = avisos.creadas[0]
    assert user_ids == [falta_1, falta_2] and tipo == "DUE_SOON" and titulo == "Cierra pronto: Informe"
    assert datos["event_at"] == cierre
    assert avisos.correos[0][0] == [falta_1]
    # Una hora después corre de nuevo: nadie recibe el aviso por segunda vez.
    assert segundo["notificaciones"] == 0 and len(avisos.correos) == 1


def test_el_recordatorio_busca_tareas_que_cierran_en_las_proximas_24_horas(app, monkeypatch):
    consultas = []

    class _Captura(ConsultaFalsa):
        def __getattr__(self, nombre):
            if nombre == "filter":
                return lambda *criterios: (consultas.extend(criterios), self)[1]
            return super().__getattr__(nombre)

    with app.app_context():
        monkeypatch.setattr(notificaciones.db.session, "query", lambda *a: _Captura(todos=[]))
        monkeypatch.setattr(notificaciones.db.session, "commit", lambda: None)
        assert notificaciones.recordar_vencimientos(AHORA) == {"tareas": 0, "notificaciones": 0}

    sql = " AND ".join(str(c.compile(dialect=postgresql.dialect(), compile_kwargs={"literal_binds": True})) for c in consultas)
    assert "course_content_items.type = 'ASSIGNMENT'" in sql
    assert "course_content_items.due_at >" in sql and "course_content_items.due_at <=" in sql


def test_fecha_en_hora_de_colombia():
    assert notificaciones.fecha_en_colombia(datetime(2026, 1, 1, 4, 30, tzinfo=timezone.utc)) == "31 dic 2025, 11:30 p. m. (hora de Colombia)"
    assert notificaciones.fecha_en_colombia(datetime(2026, 6, 15, 17, 5, tzinfo=timezone.utc)) == "15 jun 2026, 12:05 p. m. (hora de Colombia)"
    # Una fecha sin zona (columna TIMESTAMP) se toma como UTC.
    assert notificaciones.fecha_en_colombia(datetime(2026, 6, 15, 5, 0)) == "15 jun 2026, 12:00 a. m. (hora de Colombia)"


def test_correos_de_notificacion_salen_en_un_solo_pedido_por_lotes(app, monkeypatch):
    pedidos = []

    def post(url, auth=None, data=None, timeout=None):
        pedidos.append(data)
        return SimpleNamespace(raise_for_status=lambda: None)

    monkeypatch.setitem(app.config, "MAILGUN_API_KEY", "clave-de-prueba")
    monkeypatch.setitem(app.config, "MAILGUN_DOMAIN", "mail.ejemplo.test")
    with app.app_context():
        monkeypatch.setattr(mailer.requests, "post", post)
        mailer.send_notification_emails(
            [("ana@x.com", "Ana"), ("luis@x.com", "Luis"), (None, "Sin correo")], "Nueva tarea: <b>X</b>", "Texto <script>",
        )

    assert len(pedidos) == 1
    assert pedidos[0]["to"] == ["ana@x.com", "luis@x.com"]
    # Cada destinatario recibe su copia y no ve a los demás.
    assert json.loads(pedidos[0]["recipient-variables"]) == {"ana@x.com": {"nombre": "Ana"}, "luis@x.com": {"nombre": "Luis"}}
    assert "<script>" not in pedidos[0]["html"] and "&lt;script&gt;" in pedidos[0]["html"]


def test_sin_mailgun_las_notificaciones_por_correo_no_fallan(app, monkeypatch):
    monkeypatch.setitem(app.config, "MAILGUN_API_KEY", None)
    monkeypatch.setitem(app.config, "MAILGUN_DOMAIN", None)

    def prohibido(*a, **kw):
        raise AssertionError("sin Mailgun no se llama a la red")

    with app.app_context():
        monkeypatch.setattr(mailer.requests, "post", prohibido)
        mailer.send_notification_emails([("ana@x.com", "Ana")], "Asunto", "Texto")


def test_plantilla_del_correo_escapa_el_contenido():
    html = notification_email_html('Nueva tarea: "><img src=x onerror=alert(1)>', "Curso <b>Gastro</b>")
    assert "<img" not in html and "<b>Gastro" not in html
    assert "%recipient.nombre%" in html


def _notificacion(user_id, leida=False, **kw):
    datos = dict(
        id=uuid.uuid4(), user_id=user_id, type="ASSIGNMENT_PUBLISHED", title="Nueva tarea: Informe", body="En Gastro I.",
        course_id=uuid.uuid4(), entity_type="assignment", entity_id=uuid.uuid4(), created_at=AHORA,
        read_at=AHORA if leida else None,
    )
    datos.update(kw)
    return Notification(**datos)


def _preparar_notificaciones(app, monkeypatch, user, filas, encontrada=None):
    commits = []
    consulta = ConsultaFalsa(first=encontrada, todos=filas)
    consulta.count = lambda: len([n for n in filas if n.read_at is None])
    consulta.update = lambda valores, **kw: [setattr(n, "read_at", valores["read_at"]) for n in filas if n.read_at is None] and 2
    with app.app_context():
        monkeypatch.setattr(noti_routes, "get_current_user", lambda: user)
        monkeypatch.setattr(noti_routes.Notification, "query", consulta)
        monkeypatch.setattr(noti_routes.db.session, "commit", lambda: commits.append(True))
    return commits


def test_listar_notificaciones_con_leidas_y_no_leidas(app, client, monkeypatch):
    user = SimpleNamespace(id=uuid.uuid4(), role="STUDENT")
    filas = [_notificacion(user.id), _notificacion(user.id, leida=True), _notificacion(user.id)]
    _preparar_notificaciones(app, monkeypatch, user, filas)

    body = client.get("/api/notificaciones?per_page=2", headers=_headers(app, "STUDENT", user.id)).get_json()

    assert body["no_leidas"] == 2
    assert [n["read"] for n in body["notificaciones"]] == [False, True]
    assert body["per_page"] == 2 and "total" in body and "pages" in body
    primera = body["notificaciones"][0]
    assert primera["type"] == "ASSIGNMENT_PUBLISHED" and primera["entity_type"] == "assignment"
    assert client.get("/api/notificaciones/no-leidas", headers=_headers(app, "STUDENT", user.id)).get_json() == {"no_leidas": 2}


def test_marcar_una_notificacion_como_leida(app, client, monkeypatch):
    user = SimpleNamespace(id=uuid.uuid4(), role="STUDENT")
    pendiente = _notificacion(user.id)
    commits = _preparar_notificaciones(app, monkeypatch, user, [pendiente], encontrada=pendiente)

    res = client.post(f"/api/notificaciones/{pendiente.id}/leer", headers=_headers(app, "STUDENT", user.id))

    assert res.status_code == 200 and res.get_json()["read"] is True
    assert pendiente.read_at is not None and commits == [True]


def test_no_se_puede_tocar_la_notificacion_de_otro(app, client, monkeypatch):
    user = SimpleNamespace(id=uuid.uuid4(), role="STUDENT")
    filtros = []
    consulta = ConsultaFalsa(first=None)
    consulta.filter_by = lambda **kw: (filtros.append(kw), consulta)[1]
    with app.app_context():
        monkeypatch.setattr(noti_routes, "get_current_user", lambda: user)
        monkeypatch.setattr(noti_routes.Notification, "query", consulta)
    ajena = uuid.uuid4()

    res = client.post(f"/api/notificaciones/{ajena}/leer", headers=_headers(app, "STUDENT", user.id))

    assert res.status_code == 404
    assert filtros == [{"id": ajena, "user_id": user.id}]          # se busca por id Y por dueño
    assert client.post("/api/notificaciones/no-es-uuid/leer", headers=_headers(app, "STUDENT", user.id)).status_code == 404


def test_marcar_todas_como_leidas(app, client, monkeypatch):
    user = SimpleNamespace(id=uuid.uuid4(), role="STUDENT")
    filas = [_notificacion(user.id), _notificacion(user.id)]
    _preparar_notificaciones(app, monkeypatch, user, filas)

    res = client.post("/api/notificaciones/leer-todas", headers=_headers(app, "STUDENT", user.id))

    assert res.get_json() == {"marcadas": 2, "no_leidas": 0}
    assert all(n.read_at is not None for n in filas)


def test_preferencia_de_correo_apagada_por_defecto_y_se_activa(app, client, monkeypatch):
    user = SimpleNamespace(id=uuid.uuid4(), role="STUDENT", email_notifications=False)
    with app.app_context():
        monkeypatch.setattr(noti_routes, "get_current_user", lambda: user)
        monkeypatch.setattr(noti_routes.db.session, "commit", lambda: None)
    headers = _headers(app, "STUDENT", user.id)

    assert client.get("/api/notificaciones/preferencias", headers=headers).get_json() == {"email": False}
    assert client.patch("/api/notificaciones/preferencias", json={"email": True}, headers=headers).get_json() == {"email": True}
    assert user.email_notifications is True
    assert client.patch("/api/notificaciones/preferencias", json={}, headers=headers).status_code == 400


def test_notificaciones_exigen_sesion(client):
    for metodo, ruta in (("get", ""), ("get", "/no-leidas"), ("post", "/leer-todas"), ("get", "/preferencias")):
        assert getattr(client, metodo)(f"/api/notificaciones{ruta}").status_code == 401


def test_los_recordatorios_solo_se_disparan_con_el_secreto(app, client, monkeypatch):
    llamadas = []
    with app.app_context():
        monkeypatch.setattr(notificaciones, "recordar_vencimientos", lambda: llamadas.append(True) or {"tareas": 1, "notificaciones": 3})

    # Sin secreto configurado la ruta no existe, se mande lo que se mande.
    monkeypatch.setitem(app.config, "CRON_SECRET", "")
    assert client.post("/api/notificaciones/recordatorios", headers={"X-Cron-Secret": ""}).status_code == 404

    monkeypatch.setitem(app.config, "CRON_SECRET", "un-secreto-largo")
    assert client.post("/api/notificaciones/recordatorios").status_code == 404
    assert client.post("/api/notificaciones/recordatorios", headers={"X-Cron-Secret": "otro"}).status_code == 404
    assert llamadas == []

    res = client.post("/api/notificaciones/recordatorios", headers={"X-Cron-Secret": "un-secreto-largo"})
    assert res.status_code == 200 and res.get_json() == {"tareas": 1, "notificaciones": 3}


def test_publicar_tarea_aviso_y_calificar_disparan_el_aviso_despues_del_commit(app, client, monkeypatch):
    """Las rutas llaman al servicio, y lo hacen con el cambio ya confirmado."""
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I")
    orden = []
    with app.app_context():
        monkeypatch.setattr(cc_routes, "get_current_user", lambda: docente)
        monkeypatch.setattr(cc_routes.Course, "query", ConsultaFalsa(first=curso))
        monkeypatch.setattr(cc_routes.CourseBlock, "query", ConsultaFalsa(first=object()))
        monkeypatch.setattr(cc_routes.db.session, "query", lambda *a: ConsultaFalsa(first=-1))
        monkeypatch.setattr(cc_routes.db.session, "add", lambda o: None)
        monkeypatch.setattr(cc_routes.db.session, "flush", lambda: None)
        monkeypatch.setattr(cc_routes.db.session, "commit", lambda: orden.append("commit"))
        monkeypatch.setattr(notificaciones, "tarea_publicada", lambda c, item: orden.append(("tarea", c.id, item.title)) or 1)
    ruta = f"/api/cursos/{curso.id}/bloques/{uuid.uuid4()}/contenido"
    headers = _headers(app, "TEACHER", docente.id)

    assert client.post(ruta, json={"type": "ASSIGNMENT", "title": "Informe"}, headers=headers).status_code == 201
    assert orden == ["commit", ("tarea", curso.id, "Informe")]

    # Un enlace o un cuestionario no avisan: solo las tareas.
    orden.clear()
    assert client.post(ruta, json={"type": "LINK", "title": "Guía", "link_url": "https://x.com"}, headers=headers).status_code == 201
    assert orden == ["commit"]


# ═══════════════════════════ 2. Rúbrica ═══════════════════════════

def _criterio(titulo, puntos, ident=None, descripcion=None):
    return SimpleNamespace(id=ident, title=titulo, description=descripcion, max_points=puntos)


def _nota(criterio_id, puntos, comentario=None):
    return SimpleNamespace(criterion_id=criterio_id, points=puntos, comment=comentario)


RUBRICA = [
    {"id": "c1", "title": "Anamnesis", "description": None, "max_points": 4.0},
    {"id": "c2", "title": "Diagnóstico diferencial", "description": "Al menos tres", "max_points": 6.0},
]


def test_normalizar_rubrica_asigna_ids_y_suma_el_puntaje():
    rubrica = tareas.normalizar_rubrica([_criterio(" Anamnesis ", 4, ident="c1"), _criterio("Plan", 2.5, descripcion="  ")])
    assert rubrica[0] == {"id": "c1", "title": "Anamnesis", "description": None, "max_points": 4.0}
    assert rubrica[1]["id"] and rubrica[1]["id"] != "c1" and rubrica[1]["description"] is None
    assert tareas.puntaje_de_rubrica(rubrica) == 6.5
    assert tareas.normalizar_rubrica(None) == [] and tareas.puntaje_de_rubrica([]) == 0

    with pytest.raises(tareas.ReglaInvalida):
        tareas.normalizar_rubrica([_criterio("A", 1, ident="x"), _criterio("B", 1, ident="x")])


def test_calificar_con_rubrica_suma_los_criterios():
    total, detalle = tareas.calificar_con_rubrica(RUBRICA, [_nota("c2", 4.5, " Faltó uno "), _nota("c1", 4)])
    assert total == 8.5
    assert detalle == [
        {"criterion_id": "c1", "title": "Anamnesis", "max_points": 4.0, "points": 4.0, "comment": None},
        {"criterion_id": "c2", "title": "Diagnóstico diferencial", "max_points": 6.0, "points": 4.5, "comment": "Faltó uno"},
    ]


@pytest.mark.parametrize("notas,fragmento", [
    ([_nota("c1", 4)], "Falta la nota del criterio «Diagnóstico diferencial»"),
    ([_nota("c1", 4.5), _nota("c2", 6)], "vale como máximo 4 puntos"),
    ([_nota("c1", 4), _nota("c2", 6), _nota("otro", 1)], "no es de la rúbrica"),
    ([_nota("c1", 4), _nota("c1", 3), _nota("c2", 6)], "dos veces"),
])
def test_calificar_con_rubrica_rechaza_notas_invalidas(notas, fragmento):
    with pytest.raises(tareas.ReglaInvalida) as err:
        tareas.calificar_con_rubrica(RUBRICA, notas)
    assert fragmento in str(err.value)


def _preparar_docente(app, monkeypatch, docente, curso, item=None, entrega=None):
    agregados = []
    with app.app_context():
        monkeypatch.setattr(cc_routes, "get_current_user", lambda: docente)
        monkeypatch.setattr(cc_routes.Course, "query", ConsultaFalsa(first=curso))
        monkeypatch.setattr(cc_routes.CourseBlock, "query", ConsultaFalsa(first=object()))
        monkeypatch.setattr(cc_routes.CourseContentItem, "query", ConsultaFalsa(first=item))
        monkeypatch.setattr(cc_routes.AssignmentSubmission, "query", ConsultaFalsa(first=entrega))
        monkeypatch.setattr(cc_routes.db.session, "query", lambda *a: ConsultaFalsa(first=-1))
        monkeypatch.setattr(cc_routes.db.session, "add", agregados.append)
        monkeypatch.setattr(cc_routes.db.session, "flush", lambda: None)
        monkeypatch.setattr(cc_routes.db.session, "commit", lambda: None)
        monkeypatch.setattr(cc_routes, "_archivos_de_entregas", lambda ids: {})
    return agregados


def test_crear_tarea_con_rubrica_fija_el_puntaje_maximo(app, client, monkeypatch):
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I")
    _preparar_docente(app, monkeypatch, docente, curso)

    res = client.post(
        f"/api/cursos/{curso.id}/bloques/{uuid.uuid4()}/contenido",
        json={
            "type": "ASSIGNMENT", "title": "Informe de caso", "max_score": 100,
            "rubric": [
                {"title": "Anamnesis", "max_points": 4},
                {"title": "Diagnóstico diferencial", "description": "Al menos tres", "max_points": 6},
            ],
        },
        headers=_headers(app, "TEACHER", docente.id),
    )

    assert res.status_code == 201
    body = res.get_json()
    assert [(c["title"], c["max_points"]) for c in body["rubric"]] == [("Anamnesis", 4.0), ("Diagnóstico diferencial", 6.0)]
    assert all(c["id"] for c in body["rubric"])
    assert body["max_score"] == 10.0          # la suma de la rúbrica manda sobre el max_score enviado


@pytest.mark.parametrize("rubrica", [
    [{"title": "", "max_points": 4}],
    [{"title": "Anamnesis", "max_points": 0}],
    [{"title": "Anamnesis", "max_points": -1}],
    [{"title": "Anamnesis"}],
])
def test_crear_tarea_con_rubrica_invalida_da_400(app, client, monkeypatch, rubrica):
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I")
    _preparar_docente(app, monkeypatch, docente, curso)
    res = client.post(
        f"/api/cursos/{curso.id}/bloques/{uuid.uuid4()}/contenido",
        json={"type": "ASSIGNMENT", "title": "Informe", "rubric": rubrica}, headers=_headers(app, "TEACHER", docente.id),
    )
    assert res.status_code == 400


def _tarea(**kw):
    datos = dict(
        id=uuid.uuid4(), block_id=uuid.uuid4(), type="ASSIGNMENT", title="Informe", position=0, max_score=10,
        allow_late=False, allowed_extensions=[], max_files=1, rubric=None,
    )
    datos.update(kw)
    return CourseContentItem(**datos)


def test_editar_la_rubrica_y_quitarla(app, client, monkeypatch):
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I")
    item = _tarea(rubric=list(RUBRICA), max_score=10)
    _preparar_docente(app, monkeypatch, docente, curso, item=item)
    ruta = f"/api/cursos/{curso.id}/bloques/{item.block_id}/contenido/{item.id}"
    headers = _headers(app, "TEACHER", docente.id)

    res = client.patch(ruta, json={"rubric": [
        {"id": "c1", "title": "Anamnesis completa", "max_points": 5},
        {"title": "Plan de manejo", "max_points": 10},
    ]}, headers=headers)
    assert res.status_code == 200
    assert item.rubric[0] == {"id": "c1", "title": "Anamnesis completa", "description": None, "max_points": 5.0}
    assert float(item.max_score) == 15.0

    # Editar otra cosa no toca la rúbrica.
    assert client.patch(ruta, json={"title": "Informe final"}, headers=headers).status_code == 200
    assert len(item.rubric) == 2

    assert client.patch(ruta, json={"rubric": []}, headers=headers).status_code == 200
    assert item.rubric is None and float(item.max_score) == 15.0          # vuelve a ser de nota única


def test_calificar_una_entrega_con_rubrica(app, client, monkeypatch):
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I")
    item = _tarea(rubric=list(RUBRICA))
    entrega = AssignmentSubmission(id=uuid.uuid4(), item_id=item.id, student_id=uuid.uuid4(), current_version=1)
    agregados = _preparar_docente(app, monkeypatch, docente, curso, item=item, entrega=entrega)
    avisados = []
    with app.app_context():
        monkeypatch.setattr(notificaciones, "tarea_calificada", lambda c, i, e: avisados.append(float(e.score)) or 1)

    res = client.patch(
        f"/api/cursos/{curso.id}/bloques/{item.block_id}/contenido/{item.id}/entregas/{entrega.id}",
        json={
            "feedback": "Buen trabajo en general",
            "criteria": [
                {"criterion_id": "c1", "points": 3.5, "comment": "Faltaron antecedentes"},
                {"criterion_id": "c2", "points": 6},
            ],
        },
        headers=_headers(app, "TEACHER", docente.id),
    )

    assert res.status_code == 200
    body = res.get_json()
    assert body["score"] == 9.5 and body["feedback"] == "Buen trabajo en general"
    assert [(c["title"], c["points"], c["max_points"], c["comment"]) for c in body["rubric_scores"]] == [
        ("Anamnesis", 3.5, 4.0, "Faltaron antecedentes"), ("Diagnóstico diferencial", 6.0, 6.0, None),
    ]
    assert avisados == [9.5]
    auditado = next(o for o in agregados if isinstance(o, AuditLog))
    assert auditado.new_value["score"] == 9.5 and len(auditado.new_value["rubric_scores"]) == 2


@pytest.mark.parametrize("cuerpo,fragmento", [
    ({"score": 9}, "se califica con rúbrica"),
    ({"criteria": [{"criterion_id": "c1", "points": 4}]}, "Falta la nota del criterio"),
    ({"criteria": [{"criterion_id": "c1", "points": 9}, {"criterion_id": "c2", "points": 6}]}, "vale como máximo"),
    ({"criteria": [{"criterion_id": "c1", "points": -1}, {"criterion_id": "c2", "points": 6}]}, None),
])
def test_calificar_con_rubrica_mal_da_400_y_no_cambia_la_nota(app, client, monkeypatch, cuerpo, fragmento):
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I")
    item = _tarea(rubric=list(RUBRICA))
    entrega = AssignmentSubmission(id=uuid.uuid4(), item_id=item.id, student_id=uuid.uuid4(), current_version=1, score=5)
    agregados = _preparar_docente(app, monkeypatch, docente, curso, item=item, entrega=entrega)

    res = client.patch(
        f"/api/cursos/{curso.id}/bloques/{item.block_id}/contenido/{item.id}/entregas/{entrega.id}",
        json=cuerpo, headers=_headers(app, "TEACHER", docente.id),
    )

    assert res.status_code == 400
    if fragmento:
        assert fragmento in res.get_json()["message"]
    assert float(entrega.score) == 5.0 and entrega.rubric_scores is None and agregados == []


def test_tarea_sin_rubrica_se_sigue_calificando_con_nota_unica(app, client, monkeypatch):
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I")
    item = _tarea()
    entrega = AssignmentSubmission(id=uuid.uuid4(), item_id=item.id, student_id=uuid.uuid4(), current_version=1)
    _preparar_docente(app, monkeypatch, docente, curso, item=item, entrega=entrega)
    ruta = f"/api/cursos/{curso.id}/bloques/{item.block_id}/contenido/{item.id}/entregas/{entrega.id}"
    headers = _headers(app, "TEACHER", docente.id)

    assert client.patch(ruta, json={"score": 8, "feedback": "Bien"}, headers=headers).get_json()["score"] == 8.0
    assert entrega.rubric_scores is None
    assert client.patch(ruta, json={"feedback": "sin nota"}, headers=headers).status_code == 400
    assert client.patch(ruta, json={"criteria": [{"criterion_id": "c1", "points": 1}]}, headers=headers).status_code == 400


# ═══════════════════════════ 3. Exportar calificaciones ═══════════════════════════

ITEMS = [
    {"id": "t1", "type": "ASSIGNMENT", "title": "Informe de caso", "max_score": 10.0},
    {"id": "q1", "type": "QUIZ", "title": "Parcial 1", "max_score": 20.0},
    {"id": "t2", "type": "ASSIGNMENT", "title": "=HYPERLINK(\"http://malo\")", "max_score": None},
]
ESTUDIANTES = [
    {"student_code": "EST001", "nombre": "Ana Ruiz", "email": "ana@x.com", "promedio": 85.0,
     "calificaciones": {"t1": {"score": 8.5, "status": "graded"}, "q1": {"score": 17, "status": "graded"}}},
    {"student_code": "EST002", "nombre": "=cmd|' /C calc'!A0", "email": "+luis@x.com", "promedio": None,
     "calificaciones": {"t1": {"score": None, "status": "submitted"}}},
]


def test_tabla_de_exportacion():
    encabezados, filas = exportar_calificaciones.construir_tabla(ITEMS, ESTUDIANTES)
    assert encabezados == [
        "Código", "Estudiante", "Correo", "Tarea: Informe de caso (máx. 10)", "Cuestionario: Parcial 1 (máx. 20)",
        "Tarea: =HYPERLINK(\"http://malo\")", "Promedio (%)",
    ]
    assert filas[0] == ["EST001", "Ana Ruiz", "ana@x.com", 8.5, 17.0, "", 85.0]
    # Entregado sin calificar, no entregado, y celdas que una hoja tomaría por fórmula.
    assert filas[1] == ["EST002", "'=cmd|' /C calc'!A0", "'+luis@x.com", "Sin calificar", "", "", ""]


def test_exportar_a_csv():
    contenido = exportar_calificaciones.exportar("csv", ITEMS, ESTUDIANTES)
    assert contenido.startswith(b"\xef\xbb\xbf")          # BOM: Excel lee bien tildes y eñes
    filas = list(csv.reader(io.StringIO(contenido.decode("utf-8-sig"))))
    assert filas[0][:4] == ["Código", "Estudiante", "Correo", "Tarea: Informe de caso (máx. 10)"]
    assert filas[1] == ["EST001", "Ana Ruiz", "ana@x.com", "8.5", "17.0", "", "85.0"]
    assert filas[2][1] == "'=cmd|' /C calc'!A0" and filas[2][3] == "Sin calificar"


def test_exportar_a_xlsx():
    contenido = exportar_calificaciones.exportar("xlsx", ITEMS, ESTUDIANTES, titulo_hoja="Gastro: I/II [2026]")
    assert contenido[:2] == b"PK"
    hoja = load_workbook(io.BytesIO(contenido)).active
    assert hoja.title == "Gastro  I II  2026 "          # sin los caracteres que Excel no acepta en el nombre
    filas = [[c.value for c in fila] for fila in hoja.iter_rows()]
    assert filas[0][1] == "Estudiante" and filas[0][-1] == "Promedio (%)"
    assert filas[1][:5] == ["EST001", "Ana Ruiz", "ana@x.com", 8.5, 17]
    assert hoja["D2"].data_type == "n"          # las notas son números, se pueden sumar
    # Ninguna celda quedó como fórmula.
    assert all(c.data_type != "f" for fila in hoja.iter_rows() for c in fila)
    assert hoja["B3"].value == "'=cmd|' /C calc'!A0" and hoja["F1"].data_type == "s"
    assert hoja.freeze_panes == "C2"


def test_nombre_del_archivo_exportado():
    fecha = datetime(2026, 10, 11, tzinfo=timezone.utc)
    assert exportar_calificaciones.nombre_de_archivo("Gastroenterología I / 2026", "xlsx", fecha) == "calificaciones-gastroenterologia-i-2026-2026-10-11.xlsx"
    assert exportar_calificaciones.nombre_de_archivo('"; rm -rf', "csv", fecha) == "calificaciones-rm-rf-2026-10-11.csv"
    assert exportar_calificaciones.nombre_de_archivo("", "csv", fecha) == "calificaciones-curso-2026-10-11.csv"


def _preparar_exportacion(app, monkeypatch, user, curso):
    agregados = []
    matricula = (SimpleNamespace(), SimpleNamespace(student_code="EST001"),
                 SimpleNamespace(id=uuid.uuid4(), first_name="Ana", last_name="Ruiz", email="ana@x.com"))
    with app.app_context():
        monkeypatch.setattr(cal_routes, "get_current_user", lambda: user)
        monkeypatch.setattr(cal_routes.Course, "query", ConsultaFalsa(first=curso))
        monkeypatch.setattr(cal_routes.db.session, "query", lambda *a: ConsultaFalsa(todos=[matricula]))
        monkeypatch.setattr(cal_routes, "_construir_matriz", lambda course_id, ids: (ITEMS[:2], {
            str(matricula[2].id): {"t1": {"score": 8.5, "status": "graded"}},
        }))
        monkeypatch.setattr(cal_routes.db.session, "add", agregados.append)
        monkeypatch.setattr(cal_routes.db.session, "commit", lambda: None)
    return agregados


@pytest.mark.parametrize("formato,mime,firma", [
    ("csv", "text/csv", b"\xef\xbb\xbf"),
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK"),
])
def test_el_docente_dueno_descarga_el_libro(app, client, monkeypatch, formato, mime, firma):
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I")
    agregados = _preparar_exportacion(app, monkeypatch, docente, curso)

    res = client.get(f"/api/cursos/{curso.id}/calificaciones/exportar?formato={formato}", headers=_headers(app, "TEACHER", docente.id))

    assert res.status_code == 200
    assert res.mimetype == mime and res.data.startswith(firma)
    assert res.headers["Content-Disposition"].startswith('attachment; filename="calificaciones-gastro-i-')
    assert res.headers["Content-Disposition"].endswith(f'.{formato}"')
    assert res.headers["Cache-Control"] == "no-store"
    # La descarga queda auditada.
    auditado = next(o for o in agregados if isinstance(o, AuditLog))
    assert auditado.action == "GRADEBOOK_EXPORT" and auditado.new_value == {"formato": formato, "estudiantes": 1}
    if formato == "csv":
        assert "Ana Ruiz,ana@x.com,8.5" in res.data.decode("utf-8-sig")


def test_exportar_solo_los_cursos_propios(app, client, monkeypatch):
    otro_docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=uuid.uuid4(), name="Gastro I")
    agregados = _preparar_exportacion(app, monkeypatch, otro_docente, curso)
    ruta = f"/api/cursos/{curso.id}/calificaciones/exportar"

    assert client.get(ruta, headers=_headers(app, "TEACHER", otro_docente.id)).status_code == 403
    assert client.get(ruta, headers=_headers(app, "STUDENT", uuid.uuid4())).status_code == 403
    assert client.get(ruta).status_code == 401
    assert agregados == []


def test_exportar_con_formato_desconocido_da_400(app, client, monkeypatch):
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = SimpleNamespace(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I")
    _preparar_exportacion(app, monkeypatch, docente, curso)
    res = client.get(f"/api/cursos/{curso.id}/calificaciones/exportar?formato=pdf", headers=_headers(app, "TEACHER", docente.id))
    assert res.status_code == 400


# ═══════════════════════════ 4. Indicadores de la card ═══════════════════════════

def _item(course_id, titulo, due_en=None, open_en=None, allow_late=False):
    return (
        CourseContentItem(
            id=uuid.uuid4(), block_id=uuid.uuid4(), type="ASSIGNMENT", title=titulo,
            due_at=(datetime.now(timezone.utc) + due_en) if due_en is not None else None,
            open_at=(datetime.now(timezone.utc) + open_en) if open_en is not None else None,
            allow_late=allow_late,
        ),
        course_id,
    )


def _preparar_indicadores(app, monkeypatch, tareas_de_los_cursos, entregadas=()):
    consultas = []

    def consulta(*columnas):
        consultas.append(columnas)
        if columnas[0] is cursos_routes.CourseContentItem:
            return ConsultaFalsa(todos=tareas_de_los_cursos)
        return ConsultaFalsa(todos=[(i,) for i in entregadas])

    with app.app_context():
        monkeypatch.setattr(cursos_routes.db.session, "query", consulta)
    return consultas


def test_indicadores_cuentan_pendientes_y_eligen_el_cierre_mas_cercano(app, monkeypatch):
    curso_a, curso_b, curso_c = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    lejana = _item(curso_a, "Ensayo", due_en=timedelta(days=5))
    cercana = _item(curso_a, "Informe", due_en=timedelta(hours=6))
    sin_fecha = _item(curso_a, "Bitácora")
    entregada = _item(curso_a, "Ya entregada", due_en=timedelta(hours=1))
    cerrada = _item(curso_a, "Cerrada", due_en=timedelta(days=-1))
    tardia = _item(curso_b, "Acepta tardías", due_en=timedelta(days=-1), allow_late=True)
    no_abre = _item(curso_b, "Todavía no abre", due_en=timedelta(days=9), open_en=timedelta(days=2))
    consultas = _preparar_indicadores(
        app, monkeypatch, [lejana, cercana, sin_fecha, entregada, cerrada, tardia, no_abre], entregadas=[entregada[0].id],
    )

    with app.app_context():
        indicadores = cursos_routes._indicadores_del_estudiante([curso_a, curso_b, curso_c], uuid.uuid4())

    a = indicadores[str(curso_a)]
    assert a["pending_assignments"] == 3          # lejana, cercana y sin fecha
    assert a["next_due_title"] == "Informe" and a["next_due_at"].endswith("+00:00")
    assert set(a) == {"pending_assignments", "next_due_at", "next_due_title"}
    # Cerrada pero con tardías: sigue pendiente, aunque ya no hay un cierre por delante.
    assert indicadores[str(curso_b)] == {"pending_assignments": 1, "next_due_at": None, "next_due_title": None}
    assert str(curso_c) not in indicadores
    assert len(consultas) == 2          # dos consultas para TODOS los cursos, no por curso


def test_mios_del_estudiante_trae_los_indicadores_en_cada_curso(app, client, monkeypatch):
    estudiante = SimpleNamespace(id=uuid.uuid4(), role="STUDENT")
    con_tareas = Course(id=uuid.uuid4(), teacher_id=uuid.uuid4(), name="Gastro I", enrollment_mode="CODE")
    sin_tareas = Course(id=uuid.uuid4(), teacher_id=uuid.uuid4(), name="Semiología", enrollment_mode="CODE")
    cierre = (AHORA + timedelta(hours=6)).isoformat()
    with app.app_context():
        monkeypatch.setattr(cursos_routes, "get_current_user", lambda: estudiante)
        monkeypatch.setattr(cursos_routes.Course, "query", ConsultaFalsa(todos=[con_tareas, sin_tareas]))
        monkeypatch.setattr(cursos_routes.User, "query", ConsultaFalsa(todos=[]))
        monkeypatch.setattr(cursos_routes, "_indicadores_del_estudiante", lambda course_ids, student_id: {
            str(con_tareas.id): {"pending_assignments": 2, "next_due_at": cierre, "next_due_title": "Informe"},
        })

    body = client.get("/api/cursos/mios", headers=_headers(app, "STUDENT", estudiante.id)).get_json()

    assert (body[0]["pending_assignments"], body[0]["next_due_at"], body[0]["next_due_title"]) == (2, cierre, "Informe")
    assert (body[1]["pending_assignments"], body[1]["next_due_at"]) == (0, None)


def test_mios_del_docente_no_calcula_indicadores(app, client, monkeypatch):
    docente = SimpleNamespace(id=uuid.uuid4(), role="TEACHER")
    curso = Course(id=uuid.uuid4(), teacher_id=docente.id, name="Gastro I", enrollment_mode="CODE")

    def prohibido(*a, **kw):
        raise AssertionError("los indicadores son solo para el estudiante")

    with app.app_context():
        monkeypatch.setattr(cursos_routes, "get_current_user", lambda: docente)
        monkeypatch.setattr(cursos_routes.Course, "query", ConsultaFalsa(todos=[curso]))
        monkeypatch.setattr(cursos_routes, "_indicadores_del_estudiante", prohibido)

    body = client.get("/api/cursos/mios", headers=_headers(app, "TEACHER", docente.id)).get_json()
    assert "pending_assignments" not in body[0]
